"""Turning documents, links and pasted text into paragraph segments (with page numbers for PDFs)."""

from __future__ import annotations

import html
import ipaddress
import json
import re
import socket
import time
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import urlparse

import httpx2 as httpx

from isnad_worker.jobs import PermanentError
from isnad_worker.text import is_rtl

MAX_DOWNLOAD_BYTES = 30 * 1024 * 1024
# A whole download, however slowly the server trickles bytes (the read timeout alone never fires).
DOWNLOAD_DEADLINE_SECONDS = 120
USER_AGENT = "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126 Safari/537.36 Isnad/1.0"


@dataclass
class Paragraph:
    text: str
    page: int | None = None


_BLANK_LINES = re.compile(r"\n\s*\n+")
_SPACES = re.compile(r"[ \t ]+")


def paragraphs_from_text(text: str, page: int | None = None, min_chars: int = 2) -> list[Paragraph]:
    text = text.replace("\r\n", "\n")
    chunks = _BLANK_LINES.split(text) if _BLANK_LINES.search(text) else text.split("\n")
    out = []
    for chunk in chunks:
        cleaned = _SPACES.sub(" ", " ".join(line.strip() for line in chunk.splitlines())).strip()
        if len(cleaned) >= min_chars:
            out.append(Paragraph(cleaned, page))
    return out


def _pdf_blocks(page) -> list[str]:
    """A page's text blocks in reading order, each line in logical word order.

    PDF text layers keep Arabic in visual order (words left to right), so Arabic lines are
    rebuilt from word positions, right to left. Blocks keep columns apart.
    """
    lines: dict[tuple[int, int], list[tuple]] = {}
    for word in page.get_text("words"):
        lines.setdefault((word[5], word[6]), []).append(word)
    blocks: dict[int, list[list[tuple]]] = {}
    for (block, _), words in lines.items():
        blocks.setdefault(block, []).append(words)
    out: list[str] = []
    for block in sorted(blocks):
        rendered = []
        for words in sorted(blocks[block], key=lambda ws: min(w[1] for w in ws)):
            rtl = is_rtl(" ".join(w[4] for w in words))
            ordered = sorted(words, key=lambda w: -w[2] if rtl else w[0])
            rendered.append(" ".join(w[4] for w in ordered))
        text = "\n".join(rendered).strip()
        if text:
            out.append(text)
    return out


def paragraphs_from_pdf(path: Path) -> list[Paragraph]:
    import pymupdf

    out: list[Paragraph] = []
    with pymupdf.open(path) as doc:
        for index, page in enumerate(doc, start=1):
            for block in _pdf_blocks(page):
                text = _SPACES.sub(" ", block.replace("\n", " ")).strip()
                if len(text) >= 3:
                    out.append(Paragraph(text, index))
    return out


def pdf_page_texts(path: Path) -> list[str]:
    import pymupdf

    with pymupdf.open(path) as doc:
        return ["\n\n".join(_pdf_blocks(page)) for page in doc]


def paragraphs_from_docx(path: Path) -> list[Paragraph]:
    import docx

    document = docx.Document(str(path))
    out = [Paragraph(p.text.strip()) for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                out.append(Paragraph(" | ".join(dict.fromkeys(cells))))
    return out


_BLOCKS = ["h1", "h2", "h3", "p", "li", "td", "blockquote"]
_BREAK = "\u2029"  # paragraph separator standing in for <br> while reading blocks


def _link_ratio(el) -> float:
    """Share of an element's text that sits inside links (menus and related-story lists are mostly links)."""
    text = len(el.get_text(" ", strip=True)) or 1
    linked = sum(len(a.get_text(" ", strip=True)) for a in el.find_all("a"))
    return linked / text


def _densest(soup):
    """The element whose paragraphs hold the most text: each paragraph credits its parent fully and
    its grandparent by half (the classic readability scoring). None when the page has no prose."""
    scores: dict[int, float] = {}
    nodes: dict[int, object] = {}
    for p in soup.find_all("p"):
        length = len(p.get_text(" ", strip=True))
        if length < 40 or _link_ratio(p) > 0.5:
            continue
        for parent, weight in ((p.parent, 1.0), (p.parent.parent if p.parent else None, 0.5)):
            if parent is not None:
                scores[id(parent)] = scores.get(id(parent), 0) + length * weight
                nodes[id(parent)] = parent
    return nodes[max(scores, key=scores.__getitem__)] if scores else None


def _main_container(soup):
    """The element holding the article body. The largest <article> (else <main>) bounds it; inside
    that, the densest paragraph container narrows it further (dropping related-story teasers)."""

    def size(el) -> int:
        return len(_SPACES.sub(" ", el.get_text(" ", strip=True)))

    articles = soup.find_all("article")
    bound = max(articles, key=size) if articles else soup.find("main")
    if bound is not None and size(bound) < 60:
        bound = None
    densest = _densest(soup)
    if bound is None:
        return densest or soup.body or soup
    if densest is not None and (densest is bound or bound in densest.parents):
        # Only narrow when most of the bound's prose is there.
        return densest if size(densest) >= 0.5 * size(bound) else bound
    return bound


def _blocks(soup, root) -> list[Paragraph]:
    """Text blocks in reading order under `root` (plus the page headline), skipping link lists."""
    out: list[Paragraph] = []
    seen: set[str] = set()
    headline = soup.find("h1")
    if headline is not None and root is not soup and headline not in root.descendants:
        text = _SPACES.sub(" ", headline.get_text(" ", strip=True)).strip()
        if text:
            seen.add(text)
            out.append(Paragraph(text))
    # Line breaks separate paragraphs on sites that set a whole story in one block with <br>s.
    for br in root.find_all("br"):
        br.replace_with(_BREAK)
    for el in root.find_all(_BLOCKS + ["div"]):
        if el.find(_BLOCKS + ["div"]):  # a block that contains blocks: take the inner ones instead
            continue
        if _link_ratio(el) > 0.6:
            continue
        for piece in el.get_text(" ").split(_BREAK):
            text = _SPACES.sub(" ", piece).strip()
            if len(text) < 3 or text in seen:
                continue
            if el.name in ("li", "td", "div") and len(text) < 25:
                continue
            seen.add(text)
            out.append(Paragraph(text))
    return out


def text_from_html(html: str) -> tuple[str | None, list[Paragraph]]:
    """Main readable content of a web page (press release, article), without menus and boilerplate."""
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html, "lxml")
    meta_title = soup.find("meta", attrs={"property": "og:title"})
    title = (
        meta_title.get("content", "").strip() if meta_title and meta_title.get("content") else None
    ) or (soup.title.get_text(strip=True) if soup.title else None)
    # Find the content first: some government sites (the UAE design system) put <main> inside a
    # <header>, and ASP.NET sites wrap the whole page in a <form>, so boilerplate is removed only
    # where it doesn't contain the content.
    root = _main_container(soup)
    containing = {id(parent) for parent in root.parents}
    for tag in soup(
        [
            "script",
            "style",
            "noscript",
            "nav",
            "header",
            "footer",
            "aside",
            "iframe",
            "svg",
            "input",
            "select",
            "button",
            "textarea",
        ]
    ):
        if id(tag) not in containing and not tag.decomposed:
            tag.decompose()
    out = _blocks(soup, root)
    # Some sites set their text in <div>s and <span>s: when the tagged blocks miss most of the
    # content, read the content root line by line instead.
    total = len(_SPACES.sub(" ", root.get_text(" ", strip=True)))
    if sum(len(p.text) for p in out) < 0.5 * total:
        out = [
            Paragraph(line)
            for line in dict.fromkeys(
                _SPACES.sub(" ", raw).strip() for raw in root.get_text("\n").splitlines()
            )
            if len(line) >= 3
        ]
    if sum(len(p.text) for p in out) < 60 and root is not (soup.body or soup):
        out = _blocks(soup, soup.body or soup)
    return title, out


def _assert_public_host(url: str) -> None:
    """Refuse URLs that resolve to private, loopback or link-local addresses (SSRF protection)."""
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https") or not parsed.hostname:
        raise PermanentError("Only http(s) links are supported.")
    try:
        infos = socket.getaddrinfo(
            parsed.hostname, parsed.port or (443 if parsed.scheme == "https" else 80)
        )
    except socket.gaierror as exc:
        raise PermanentError(f"The link's host couldn't be resolved: {parsed.hostname}") from exc
    for info in infos:
        address = ipaddress.ip_address(info[4][0])
        if (
            address.is_private
            or address.is_loopback
            or address.is_link_local
            or address.is_reserved
            or address.is_multicast
        ):
            raise PermanentError("That link points to a private network address.")


@dataclass
class Download:
    url: str
    content: bytes
    content_type: str


# Sites whose article pages are a script shell (the text is loaded by the browser): their articles
# are read through the site's own article API instead.
_WAM_ARTICLE = re.compile(r"^/(?:ar|en)/article/([^/?#]+)")
WAM_ARTICLE_API = "https://www.wam.ae/api/app/articles/GetArticleBySlug?slug={slug}"


def wam_article_html(raw: bytes) -> tuple[str, str | None]:
    """wam.ae's article API JSON as a plain HTML article, and its date."""
    article = json.loads(raw.decode("utf-8"))
    title = (article.get("title") or "").strip()
    date = (article.get("articleDate") or "")[:10] or None
    body = article.get("body") or ""
    page = (
        "<!doctype html><html><head><meta charset='utf-8'>"
        f"<title>{html.escape(title)}</title></head><body><article>"
        f"<h1>{html.escape(title)}</h1>"
        + (f"<p>{html.escape(date)}</p>" if date else "")
        + body
        + "</article></body></html>"
    )
    return page, date


def article_api(url: str) -> str | None:
    """The API URL that returns this article's text, for sites that render articles in the browser."""
    parsed = urlparse(url)
    host = (parsed.hostname or "").lower().removeprefix("www.")
    m = _WAM_ARTICLE.match(parsed.path)
    if host == "wam.ae" and m:
        return WAM_ARTICLE_API.format(slug=m.group(1))
    return None


def download(url: str, *, max_bytes: int = MAX_DOWNLOAD_BYTES) -> Download:
    """Fetch a public URL with size and redirect limits, re-checking every redirect target. A WAM
    article comes back as readable HTML from its article API, under its page URL."""
    api = article_api(url)
    if api:
        fetched = _fetch(api, max_bytes=max_bytes)
        try:
            page, _ = wam_article_html(fetched.content)
        except (ValueError, AttributeError) as exc:
            raise PermanentError("The article couldn't be read from WAM.") from exc
        return Download(url, page.encode("utf-8"), "text/html")
    return _fetch(url, max_bytes=max_bytes)


def _fetch(url: str, *, max_bytes: int) -> Download:
    current = url
    deadline = time.monotonic() + DOWNLOAD_DEADLINE_SECONDS
    with httpx.Client(
        timeout=45,
        follow_redirects=False,
        headers={"user-agent": USER_AGENT, "accept-language": "ar,en;q=0.8"},
    ) as client:
        for _ in range(6):
            _assert_public_host(current)
            with client.stream("GET", current) as response:
                if (
                    response.status_code in (301, 302, 303, 307, 308)
                    and "location" in response.headers
                ):
                    current = str(response.url.join(response.headers["location"]))
                    continue
                if response.status_code >= 400:
                    raise PermanentError(f"The link returned HTTP {response.status_code}.")
                chunks: list[bytes] = []
                size = 0
                for chunk in response.iter_bytes(1 << 16):
                    if time.monotonic() > deadline:
                        raise PermanentError("The link took too long to download.")
                    size += len(chunk)
                    if size > max_bytes:
                        raise PermanentError("The linked file is too large.")
                    chunks.append(chunk)
                content_type = (
                    response.headers.get("content-type", "").split(";")[0].strip().lower()
                )
                return Download(str(response.url), b"".join(chunks), content_type)
    raise PermanentError("Too many redirects.")


def decode_html(data: bytes, content_type_header: str = "") -> str:
    for encoding in ("utf-8", "windows-1256", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def kind_for(content_type: str, url_or_name: str) -> str:
    name = url_or_name.lower().split("?")[0]
    if content_type == "application/pdf" or name.endswith(".pdf"):
        return "pdf"
    if "spreadsheetml" in content_type or name.endswith((".xlsx", ".xlsm")):
        return "xlsx"
    if content_type == "application/vnd.ms-excel" or name.endswith(".xls"):
        return "xls"
    if content_type == "text/csv" or name.endswith(".csv"):
        return "csv"
    if "wordprocessingml" in content_type or name.endswith(".docx"):
        return "docx"
    if content_type.startswith("text/plain") or name.endswith(".txt"):
        return "text"
    return "html"
