"""Storing official documents as immutable, versioned snapshots, and reading their text and tables."""

from __future__ import annotations

import csv
import hashlib
import io
import logging
import re
import tempfile
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import urlparse

from isnad_worker import db, storage
from isnad_worker.pipeline import content as content_mod
from isnad_worker.text import normalize_for_match

log = logging.getLogger("isnad.registry")

MIME = {
    "pdf": "application/pdf",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "xls": "application/vnd.ms-excel",
    "csv": "text/csv",
    "html": "text/html",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "text": "text/plain",
}
EXT = {
    "pdf": "pdf",
    "xlsx": "xlsx",
    "xls": "xls",
    "csv": "csv",
    "html": "html",
    "docx": "docx",
    "text": "txt",
}
# Kinds compared by their extracted text rather than their bytes (see store_version).
TEXT_KINDS = {"html", "text"}
# The documents.kind enum has no separate xls value; old Excel files are stored as xlsx-kind documents.
DB_KIND = {"xls": "xlsx"}


@dataclass
class Extracted:
    text: str
    tables: list[dict]
    page_count: int | None


def _column_letter(index: int) -> str:
    letters = ""
    index += 1
    while index:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


# Cells with no display format ("General") keep at most this many decimals.
GENERAL_DECIMALS = 6
_QUOTED = re.compile(r'"[^"]*"|\[[^\]]*\]|\\.')


def display_format(number_format: str | None) -> tuple[int | None, bool]:
    """Decimals and percent shown by an Excel number format ("#,##0.0" → 1; "0%" → 0, percent).

    Returns (None, False) for "General" and formats without a numeric pattern.
    """
    section = _QUOTED.sub("", (number_format or "General").split(";")[0])
    if "0" not in section and "#" not in section:
        return None, False
    percent = "%" in section
    decimals = 0
    if "." in section:
        decimals = len(re.match(r"[0#]*", section.split(".", 1)[1]).group(0))
    return decimals, percent


def published_number(value: float, number_format: str | None) -> tuple[float | int, int, bool]:
    """A cell's number as the sheet shows it: percent cells ×100, rounded to the shown decimals.

    Spreadsheets keep full float precision (325682.14963061886) behind a format that shows
    325,682.1; the registry stores what was published.
    """
    decimals, percent = display_format(number_format)
    shown = value * 100 if percent else value
    if decimals is None:
        rounded = round(shown, GENERAL_DECIMALS)
        return (
            (int(rounded) if float(rounded).is_integer() else rounded),
            decimals_of_float(rounded),
            percent,
        )
    rounded = round(shown, decimals)
    return (int(rounded) if decimals == 0 else rounded), decimals, percent


def decimals_of_float(value: float) -> int:
    text = repr(float(value))
    return 0 if text.endswith(".0") or "e" in text else len(text.split(".", 1)[1])


def _sheet_cells_xlsx(data: bytes) -> list[dict]:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    tables = []
    for sheet in workbook.worksheets:
        cells: dict[str, object] = {}
        decimals: dict[str, int] = {}
        percents: list[str] = []
        for r, row in enumerate(sheet.iter_rows()):
            for c, cell in enumerate(row):
                value = getattr(cell, "value", None)
                if value is None or str(value).strip() == "":
                    continue
                ref = f"{_column_letter(c)}{r + 1}"
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    number, places, percent = published_number(value, cell.number_format)
                    cells[ref] = number
                    decimals[ref] = places
                    if percent:
                        percents.append(ref)
                else:
                    cells[ref] = str(value).strip()
            if r > 5000:
                break
        if cells:
            tables.append(
                {
                    "kind": "sheet",
                    "sheet": sheet.title,
                    "cells": cells,
                    "decimals": decimals,
                    "percent": percents,
                }
            )
    workbook.close()
    return tables


def _sheet_cells_xls(data: bytes) -> list[dict]:
    import xlrd

    book = xlrd.open_workbook(file_contents=data)
    tables = []
    for sheet in book.sheets():
        cells: dict[str, object] = {}
        for r in range(min(sheet.nrows, 5000)):
            for c in range(sheet.ncols):
                value = sheet.cell_value(r, c)
                if value not in ("", None):
                    cells[f"{_column_letter(c)}{r + 1}"] = (
                        published_number(value, None)[0]
                        if isinstance(value, (int, float))
                        else str(value).strip()
                    )
        if cells:
            tables.append({"kind": "sheet", "sheet": sheet.name, "cells": cells})
    return tables


def _sheet_cells_csv(data: bytes) -> list[dict]:
    text = content_mod.decode_html(data)
    cells: dict[str, object] = {}
    for r, row in enumerate(csv.reader(io.StringIO(text))):
        for c, value in enumerate(row):
            if value.strip():
                cells[f"{_column_letter(c)}{r + 1}"] = value.strip()
        if r > 5000:
            break
    return [{"kind": "sheet", "sheet": "csv", "cells": cells}]


def _sheets_as_text(tables: list[dict]) -> str:
    lines = []
    for table in tables:
        lines.append(f"## {table['sheet']}")
        rows: dict[int, list[str]] = {}
        percent = set(table.get("percent") or [])
        for ref, value in table["cells"].items():
            row = int("".join(ch for ch in ref if ch.isdigit()))
            shown = f"{value}%" if ref in percent else value
            rows.setdefault(row, []).append(f"{ref}={shown}")
        for row in sorted(rows):
            lines.append(" | ".join(rows[row]))
    return "\n".join(lines)


def extract(kind: str, data: bytes) -> Extracted:
    if kind == "pdf":
        with tempfile.NamedTemporaryFile(suffix=".pdf") as tmp:
            tmp.write(data)
            tmp.flush()
            pages = content_mod.pdf_page_texts(Path(tmp.name))
        return Extracted("\n\n".join(pages), [{"kind": "pdf_pages", "pages": pages}], len(pages))
    if kind in ("xlsx", "xls", "csv"):
        tables = {"xlsx": _sheet_cells_xlsx, "xls": _sheet_cells_xls, "csv": _sheet_cells_csv}[
            kind
        ](data)
        return Extracted(_sheets_as_text(tables), tables, None)
    if kind == "docx":
        with tempfile.NamedTemporaryFile(suffix=".docx") as tmp:
            tmp.write(data)
            tmp.flush()
            paragraphs = content_mod.paragraphs_from_docx(Path(tmp.name))
        return Extracted("\n\n".join(p.text for p in paragraphs), [], None)
    if kind == "text":
        return Extracted(content_mod.decode_html(data), [], None)
    _, paragraphs = content_mod.text_from_html(content_mod.decode_html(data))
    return Extracted("\n\n".join(p.text for p in paragraphs), [], None)


def source_for_url(url: str, newsroom_id: str | None = None) -> dict | None:
    """The trusted source whose domains cover this URL's host."""
    host = (urlparse(url).hostname or "").lower()
    rows = db.fetch_all(
        """
        select id, slug, entity_id, newsroom_id, tier, domains from public.sources
        where status = 'active' and (newsroom_id is null or newsroom_id = %s)
        """,
        (newsroom_id,),
    )
    best = None
    for row in rows:
        for domain in row["domains"]:
            d = domain.lower().lstrip(".")
            if (host == d or host.endswith("." + d)) and (
                best is None or row["tier"] < best["tier"]
            ):
                best = row
    return best


def store_version(
    *,
    source_id: str,
    newsroom_id: str | None,
    title: str,
    url: str | None,
    kind: str,
    data: bytes,
    fetched_via: str,
    fetched_by: str | None = None,
    published_at: str | None = None,
    language: str = "ar",
    extract_now: bool = True,
) -> dict:
    """Store a document snapshot. Returns the version row (an existing one if the content is unchanged).

    Web pages change bytes on every fetch (tokens, timestamps, ads) without changing what they
    say, so pages are compared by their extracted text; files by their bytes.
    """
    extracted = extract(kind, data) if kind in TEXT_KINDS else None
    text_key = normalize_for_match(extracted.text) if extracted else None
    content_hash = (
        "text:" + hashlib.sha256(text_key.encode()).hexdigest()
        if text_key is not None
        else hashlib.sha256(data).hexdigest()
    )
    with db.transaction() as conn:
        document = None
        if url:
            document = conn.execute(
                "select * from public.documents where source_id = %s and url = %s", (source_id, url)
            ).fetchone()
        if document is None:
            document = conn.execute(
                """
                insert into public.documents (source_id, newsroom_id, title, url, kind, language, published_at)
                values (%s, %s, %s, %s, %s, %s, %s)
                returning *
                """,
                (
                    source_id,
                    newsroom_id,
                    title[:500],
                    url,
                    DB_KIND.get(kind, kind),
                    language,
                    published_at,
                ),
            ).fetchone()
        existing = conn.execute(
            "select * from public.document_versions where document_id = %s and content_hash = %s",
            (document["id"], content_hash),
        ).fetchone()
        if existing:
            return existing
        previous = conn.execute(
            "select * from public.document_versions where document_id = %s order by version desc limit 1",
            (document["id"],),
        ).fetchone()
        # Pages stored before text hashing: the same text as the latest version is no change.
        if previous and text_key is not None and previous["text_normalized"] == text_key:
            return previous
        version = (previous["version"] + 1) if previous else 1
        path = f"{source_id}/{document['id']}/v{version}.{EXT.get(kind, 'bin')}"
        storage.upload(storage.DOCUMENTS, path, data, MIME.get(kind))
        row = conn.execute(
            """
            insert into public.document_versions
              (document_id, version, storage_path, content_hash, mime_type, size_bytes, fetched_via, fetched_by,
               previous_version_id, extraction_status)
            values (%s, %s, %s, %s, %s, %s, %s, %s, %s, 'pending')
            returning *
            """,
            (
                document["id"],
                version,
                path,
                content_hash,
                MIME.get(kind, "application/octet-stream"),
                len(data),
                fetched_via,
                fetched_by,
                previous["id"] if previous else None,
            ),
        ).fetchone()
        conn.execute(
            "update public.documents set current_version_id = %s, title = %s where id = %s",
            (row["id"], title[:500] or document["title"], document["id"]),
        )
    if extract_now:
        apply_text(row["id"], kind, data, extracted)
        row = db.fetch_one("select * from public.document_versions where id = %s", (row["id"],))
    return row


def apply_text(
    version_id: str, kind: str, data: bytes, extracted: Extracted | None = None
) -> Extracted:
    extracted = extracted or extract(kind, data)
    db.execute(
        """
        update public.document_versions
        set text_content = %s, text_normalized = %s, tables = %s, page_count = %s
        where id = %s
        """,
        (
            extracted.text,
            normalize_for_match(extracted.text),
            db.json(extracted.tables),
            extracted.page_count,
            version_id,
        ),
    )
    return extracted
